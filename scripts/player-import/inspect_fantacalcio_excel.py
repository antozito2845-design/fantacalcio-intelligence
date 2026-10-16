from pathlib import Path

from openpyxl import load_workbook


excel_path = Path(
    "scripts/player-import/input/fantacalcio/"
    "Quotazioni_Fantacalcio_Stagione_2026_27.xlsx"
)

workbook = load_workbook(
    excel_path,
    read_only=True,
    data_only=True,
)

print(f"FILE: {excel_path}")
print(f"FOGLI: {workbook.sheetnames}")

for sheet_name in workbook.sheetnames:
    worksheet = workbook[sheet_name]

    print()
    print("=" * 100)
    print(
        f"FOGLIO: {sheet_name} | "
        f"RIGHE: {worksheet.max_row} | "
        f"COLONNE: {worksheet.max_column}"
    )
    print("=" * 100)

    for row_number, row in enumerate(
        worksheet.iter_rows(values_only=True),
        start=1,
    ):
        print(row_number, row)

        if row_number >= 20:
            break