from __future__ import annotations

import csv
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]

INPUT_DIRECTORY = (
    PROJECT_ROOT
    / "scripts"
    / "player-import"
    / "input"
    / "fantacalcio"
)

OUTPUT_PATH = (
    PROJECT_ROOT
    / "scripts"
    / "player-import"
    / "output"
    / "players-fantacalcio-2026-2027.csv"
)

REPORT_PATH = (
    PROJECT_ROOT
    / "scripts"
    / "player-import"
    / "logs"
    / "fantacalcio-normalization-report.txt"
)

EXPECTED_SHEET = "Tutti"

EXPECTED_HEADERS = (
    "Id",
    "R",
    "RM",
    "Nome",
    "Squadra",
    "Qt.A",
    "Qt.I",
    "Diff.",
    "Qt.A M",
    "Qt.I M",
    "Diff.M",
    "FVM",
    "FVM M",
)

VALID_ROLES = {
    "P",
    "D",
    "C",
    "A",
}

CLUB_ALIASES = {
    "Atalanta": "Atalanta",
    "Bologna": "Bologna",
    "Cagliari": "Cagliari",
    "Como": "Como",
    "Fiorentina": "Fiorentina",
    "Frosinone": "Frosinone",
    "Genoa": "Genoa",
    "Inter": "Inter",
    "Inter": "Inter",
    "Juventus": "Juventus",
    "Lazio": "Lazio",
    "Lecce": "Lecce",
    "Milan": "Milan",
    "Monza": "Monza",
    "Napoli": "Napoli",
    "Parma": "Parma",
    "Roma": "Roma",
    "Sassuolo": "Sassuolo",
    "Torino": "Torino",
    "Udinese": "Udinese",
    "Venezia": "Venezia",
}


@dataclass(frozen=True)
class PlayerRow:
    external_id: str
    name: str
    club: str
    role: str
    quotation: int
    fmv: int


def clean_text(value: object) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value).strip().split()
    )


def clean_integer(
    value: object,
    field_name: str,
    excel_row: int,
) -> int:
    if value is None or value == "":
        raise ValueError(
            f"Riga Excel {excel_row}: "
            f"{field_name} mancante"
        )

    if isinstance(value, bool):
        raise ValueError(
            f"Riga Excel {excel_row}: "
            f"{field_name} non valido: {value}"
        )

    try:
        numeric_value = int(value)
    except (TypeError, ValueError) as exception:
        raise ValueError(
            f"Riga Excel {excel_row}: "
            f"{field_name} non numerico: {value}"
        ) from exception

    if numeric_value < 0:
        raise ValueError(
            f"Riga Excel {excel_row}: "
            f"{field_name} negativo: {numeric_value}"
        )

    return numeric_value


def find_latest_excel() -> Path:
    excel_files = sorted(
        [
            path
            for path in INPUT_DIRECTORY.iterdir()
            if path.is_file()
            and path.suffix.casefold()
            in {
                ".xlsx",
                ".xlsm",
                ".xltx",
                ".xltm",
            }
        ],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )

    if not excel_files:
        raise FileNotFoundError(
            "Nessun file Excel Fantacalcio trovato in "
            f"{INPUT_DIRECTORY}"
        )

    return excel_files[0]


def validate_title(value: object) -> None:
    title = clean_text(value).casefold()

    required_fragments = (
        "quotazioni fantacalcio",
        "2026",
        "27",
    )

    missing = [
        fragment
        for fragment in required_fragments
        if fragment not in title
    ]

    if missing:
        raise RuntimeError(
            "Il file non sembra appartenere alla stagione "
            "2026/27. Titolo trovato: "
            f"{clean_text(value)}"
        )


def validate_headers(
    worksheet,
) -> dict[str, int]:
    raw_headers = next(
        worksheet.iter_rows(
            min_row=2,
            max_row=2,
            values_only=True,
        )
    )

    headers = tuple(
        clean_text(value)
        for value in raw_headers
    )

    if headers != EXPECTED_HEADERS:
        raise RuntimeError(
            "Intestazioni Excel inattese.\n"
            f"Attese: {EXPECTED_HEADERS}\n"
            f"Trovate: {headers}"
        )

    return {
        header: index
        for index, header in enumerate(headers)
    }


def normalize_club(
    source_club: str,
    excel_row: int,
) -> str:
    if source_club not in CLUB_ALIASES:
        raise ValueError(
            f"Riga Excel {excel_row}: "
            f"squadra non riconosciuta: {source_club}"
        )

    return CLUB_ALIASES[source_club]


def parse_players(
    worksheet,
    header_indexes: dict[str, int],
) -> tuple[list[PlayerRow], list[str]]:
    players: list[PlayerRow] = []
    warnings: list[str] = []
    external_ids: set[str] = set()
    player_occurrences: Counter[
        tuple[str, str, str]
    ] = Counter()

    for excel_row, values in enumerate(
        worksheet.iter_rows(
            min_row=3,
            values_only=True,
        ),
        start=3,
    ):
        if all(value is None for value in values):
            continue

        raw_id = values[
            header_indexes["Id"]
        ]

        raw_role = values[
            header_indexes["R"]
        ]

        raw_name = values[
            header_indexes["Nome"]
        ]

        raw_club = values[
            header_indexes["Squadra"]
        ]

        raw_quotation = values[
            header_indexes["Qt.A"]
        ]

        raw_fmv = values[
            header_indexes["FVM"]
        ]

        external_numeric_id = clean_integer(
            raw_id,
            "Id",
            excel_row,
        )

        role = clean_text(
            raw_role
        ).upper()

        name = clean_text(
            raw_name
        )

        source_club = clean_text(
            raw_club
        )

        quotation = clean_integer(
            raw_quotation,
            "Qt.A",
            excel_row,
        )

        fmv = clean_integer(
            raw_fmv,
            "FVM",
            excel_row,
        )

        if role not in VALID_ROLES:
            raise ValueError(
                f"Riga Excel {excel_row}: "
                f"ruolo non valido: {role}"
            )

        if not name:
            raise ValueError(
                f"Riga Excel {excel_row}: "
                "nome mancante"
            )

        club = normalize_club(
            source_club,
            excel_row,
        )

        external_id = (
            f"FANTACALCIO-{external_numeric_id}"
        )

        if external_id in external_ids:
            raise ValueError(
                f"Riga Excel {excel_row}: "
                f"ID duplicato: {external_id}"
            )

        external_ids.add(
            external_id
        )

        player_key = (
            name.casefold(),
            club.casefold(),
            role,
        )

        player_occurrences[
            player_key
        ] += 1

        players.append(
            PlayerRow(
                external_id=external_id,
                name=name,
                club=club,
                role=role,
                quotation=quotation,
                fmv=fmv,
            )
        )

    for key, count in player_occurrences.items():
        if count <= 1:
            continue

        name, club, role = key

        warnings.append(
            "Possibile omonimia o duplicato: "
            f"{name} - {club} - {role} "
            f"({count} occorrenze)"
        )

    return players, warnings


def validate_players(
    players: list[PlayerRow],
) -> list[str]:
    errors: list[str] = []

    if not players:
        errors.append(
            "Nessun giocatore estratto"
        )

    role_counts = Counter(
        player.role
        for player in players
    )

    missing_roles = sorted(
        VALID_ROLES - set(role_counts)
    )

    if missing_roles:
        errors.append(
            "Ruoli mancanti: "
            + ", ".join(missing_roles)
        )

    club_counts = Counter(
        player.club
        for player in players
    )

    expected_clubs = set(
        CLUB_ALIASES.values()
    )

    missing_clubs = sorted(
        expected_clubs - set(club_counts)
    )

    if missing_clubs:
        errors.append(
            "Squadre mancanti: "
            + ", ".join(missing_clubs)
        )

    return errors


def write_csv(
    players: list[PlayerRow],
) -> None:
    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with OUTPUT_PATH.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as csv_file:
        writer = csv.DictWriter(
            csv_file,
            fieldnames=[
                "external_id",
                "name",
                "club",
                "role",
                "quotation",
                "fmv",
            ],
        )

        writer.writeheader()

        for player in players:
            writer.writerow(
                {
                    "external_id": player.external_id,
                    "name": player.name,
                    "club": player.club,
                    "role": player.role,
                    "quotation": player.quotation,
                    "fmv": player.fmv,
                }
            )


def write_report(
    excel_path: Path,
    players: list[PlayerRow],
    warnings: list[str],
    errors: list[str],
) -> None:
    REPORT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    role_counts = Counter(
        player.role
        for player in players
    )

    club_counts = Counter(
        player.club
        for player in players
    )

    lines = [
        "FANTACALCIO 2026/2027 - REPORT NORMALIZZAZIONE",
        "=" * 60,
        f"Excel: {excel_path}",
        f"CSV: {OUTPUT_PATH}",
        f"Giocatori totali: {len(players)}",
        "",
        "CONTEGGIO PER RUOLO",
    ]

    for role in ("P", "D", "C", "A"):
        lines.append(
            f"{role}: {role_counts.get(role, 0)}"
        )

    lines.extend(
        [
            "",
            "CONTEGGIO PER SQUADRA",
        ]
    )

    for club in sorted(club_counts):
        lines.append(
            f"{club}: {club_counts[club]}"
        )

    lines.extend(
        [
            "",
            "AVVISI",
        ]
    )

    if warnings:
        lines.extend(warnings)
    else:
        lines.append(
            "Nessun avviso"
        )

    lines.extend(
        [
            "",
            "ERRORI",
        ]
    )

    if errors:
        lines.extend(errors)
    else:
        lines.append(
            "Nessun errore"
        )

    REPORT_PATH.write_text(
        "\n".join(lines) + "\n",
        encoding="utf-8",
    )


def main() -> int:
    try:
        excel_path = find_latest_excel()

        print(f"Excel: {excel_path}")

        workbook = load_workbook(
            excel_path,
            read_only=True,
            data_only=True,
        )

        if EXPECTED_SHEET not in workbook.sheetnames:
            raise RuntimeError(
                f"Foglio {EXPECTED_SHEET} non trovato. "
                f"Fogli disponibili: {workbook.sheetnames}"
            )

        worksheet = workbook[
            EXPECTED_SHEET
        ]

        validate_title(
            worksheet.cell(
                row=1,
                column=1,
            ).value
        )

        header_indexes = validate_headers(
            worksheet
        )

        players, warnings = parse_players(
            worksheet,
            header_indexes,
        )

        errors = validate_players(
            players
        )

        write_report(
            excel_path,
            players,
            warnings,
            errors,
        )

        if errors:
            print()
            print("NORMALIZZAZIONE FALLITA")

            for error in errors:
                print(f"- {error}")

            print(f"Report: {REPORT_PATH}")
            return 1

        write_csv(
            players
        )

        print()
        print("NORMALIZZAZIONE COMPLETATA")
        print(f"Giocatori: {len(players)}")
        print(f"CSV: {OUTPUT_PATH}")
        print(f"Report: {REPORT_PATH}")
        print(f"Avvisi: {len(warnings)}")

        for warning in warnings:
            print(f"- {warning}")

        role_counts = Counter(
            player.role
            for player in players
        )

        for role in ("P", "D", "C", "A"):
            print(
                f"Ruolo {role}: "
                f"{role_counts.get(role, 0)}"
            )

        return 0

    except Exception as exception:
        print(
            f"ERRORE: {exception}",
            file=sys.stderr,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())